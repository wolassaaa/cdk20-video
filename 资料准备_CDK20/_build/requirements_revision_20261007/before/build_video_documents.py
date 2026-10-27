from pathlib import Path
import json,csv,shutil,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
BASE=Path(__file__).resolve().parent.parent;OUT=BASE/'视频内容稿'
sys.path.insert(0,str(BASE/'_build/script_dependencies'))
from opencc import OpenCC
cc=OpenCC('s2t')
def tc(s):return cc.convert(s).replace('爲','為').replace('纔','才').replace('覈','核')
data=json.loads((OUT/'双语内容主稿.json').read_text(encoding='utf-8'))
shutil.copy2(BASE/'证据清单.csv',OUT/'证据索引.csv')
refs=[
('S01','National Cancer Institute. Liver and Bile Duct Cancer.','https://www.cancer.gov/types/liver','疾病定义；官方网页。'),
('S02','Ren F, et al. AlphaFold accelerates artificial intelligence powered drug discovery: efficient discovery of a novel CDK20 small molecule inhibitor. Chemical Science. 2023;14:1443–1452.','https://doi.org/10.1039/D2SC05709C','主案例；正式全文与16页化学补充材料已核查。流程、数字及实验以此为准。'),
('S03','Wang L, et al. Advances of Artificial Intelligence in Anti-Cancer Drug Design: A Review of the Past Decade. Pharmaceuticals. 2023;16:253.','https://doi.org/10.3390/ph16020253','背景综述；全文已核查，不替代主案例实验。'),
('S04','Jumper J, et al. Highly accurate protein structure prediction with AlphaFold. Nature. 2021;596:583–589.','https://doi.org/10.1038/s41586-021-03819-2','结构预测方法；全文已核查。'),
('S05','Ivanenkov YA, et al. Chemistry42: An AI-Driven Platform for Molecular Design and Optimization. Journal of Chemical Information and Modeling. 2023;63:695–701.','https://doi.org/10.1021/acs.jcim.2c01191','平台方法；全文已核查，与主案例团队有重叠，不是独立重复验证。'),
('S07','Bender A, et al. Artificial intelligence in drug discovery — what it is, where we stand and the path forward. Nature Reviews Drug Discovery. Published online 7 August 2026.','https://doi.org/10.1038/s41573-026-01496-2','评价视角；只使用公开摘要与元数据，不引用未读全文的数值。'),
('S08','Feng H, et al. Cell cycle-related kinase is a direct androgen receptor-regulated gene that drives β-catenin/T cell factor-dependent hepatocarcinogenesis. Journal of Clinical Investigation. 2011;121:3159–3175.','https://www.jci.org/articles/view/45967','前期靶点依据；全文已核查，不是2023年候选分子的动物试验。'),
('S10','Lengauer T, ed. Bioinformatics — From Genomes to Therapies. Vol. 1. Wiley-VCH, 2007. Public Chapter 1, pp.1–24; especially p.7.','https://application.wiley-vch.de/books/sample/3527312781_c01.pdf','实际读到的学术书章节；解释靶点、分子及选择性，不用于推导当前AI效率。'),
]
def runfont(r,size=11,bold=False,color='000000'):
 r.font.name='Calibri';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color);r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
def para(d,text,size=11,bold=False,color='000000',keep=False,after=5):
 p=d.add_paragraph();p.paragraph_format.keep_with_next=keep;p.paragraph_format.widow_control=True;p.paragraph_format.space_after=Pt(after);runfont(p.add_run(text),size,bold,color);return p
def hyperlink(d,text,url):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);c=OxmlElement('w:color');c.set(qn('w:val'),'254F6B');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);link.append(r);p._p.append(link)
for edition,field in [('简中','zh_sc'),('繁中','zh_tc')]:
 c=lambda s:tc(s) if edition=='繁中' else s
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(1.9);sec.footer_distance=Cm(.8)
 for name,size in [('Normal',11),('Title',21),('Subtitle',10),('Heading 1',15)]:
  st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei');st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.13
  for e in list(st._element.xpath('.//w:pBdr')):e.getparent().remove(e)
  if name in ('Title','Heading 1'):st.font.bold=True
 d.core_properties.title=c(data['title_sc'])+' '+c('英文与中文对照视频内容稿');d.core_properties.author='SCIE1007 project group'
 d.add_paragraph(c(data['title_sc']),style='Title')
 d.add_paragraph('Can AI help design a molecule worth testing?',style='Subtitle')
 para(d,c('英文＋'+('简体中文' if edition=='简中' else '繁体中文')+'对照内容稿  |  2026年10月7日'),9,color='444444')
 para(d,c('英文为口播正文，中文为逐段理解对照。段落编号、证据标注与文末资料不朗读。'),10,after=6)
 para(d,c(f'英文正文 {data["english_word_count"]:,} 词；按每分钟115–130词估算约16–18分钟，最终以只读英文的完整试讲计时为准。'),9,color='444444',after=10)
 for i,s in enumerate(data['sections']):
  head=d.add_heading(s['id']+'  '+c(s['title_sc']),1)
  head.paragraph_format.space_before=Pt(12)
  para(d,s['title_en'],10,color='444444',keep=True,after=8)
  for p in s['paragraphs']:
   para(d,p['id']+'  EN',8.5,bold=True,color='444444',keep=True,after=3)
   para(d,p['en'],11,keep=True,after=5)
   para(d,p[field],10.5,color='333333',keep=True,after=4)
   para(d,c('核对依据：')+' · '.join(p['refs']),8,color='555555',after=11)
 h=d.add_heading(c('证据与参考资料'),1);h.paragraph_format.page_break_before=True
 para(d,c('本页不属于口播。E编号对应同文件夹的《证据索引.csv》，其中记录原文章节、图号、结果与推论边界。未来比较方案是我们的建议，不是已经实施的实验。'),10,after=9)
 for sid,bib,url,note in refs:
  para(d,sid+'  '+bib,9.5,keep=True,after=3);hyperlink(d,url,url);para(d,c(note),9,color='333333',after=8)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;runfont(footer.add_run(c('视频内容稿  |  ')),8,color='555555');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
 path=OUT/f'视频内容稿_英文加{edition}.docx';d.save(path);reopened=Document(path);assert len(reopened.paragraphs)>100;print('DOCX_CREATED',path)
english='\n\n'.join(p['en'] for s in data['sections'] for p in s['paragraphs'])
(OUT/'英文口播纯文本.txt').write_text(english,encoding='utf-8')
(OUT/'阅读说明.txt').write_text('两份稿件共用同一份英文口播。选择简中或繁中对照版阅读即可。\n英文口播纯文本.txt只含应朗读的英文正文。正文2105词，时长为语速估算，未进行真实录音计时。\n证据索引.csv为原有35条证据卡，E编号与正文对应。参考资料和证据标注不朗读。\n本轮仅为内容稿，没有动画、分镜、拍摄安排、成员分工或已定时字幕。\n',encoding='utf-8')
