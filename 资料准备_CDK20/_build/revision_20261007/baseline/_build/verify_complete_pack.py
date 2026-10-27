from pathlib import Path
import csv, hashlib, json, re, sys, xml.etree.ElementTree as ET
from zipfile import ZipFile
from docx import Document
from pypdf import PdfReader
from PIL import Image
import pdfplumber
BASE=Path(__file__).resolve().parent.parent
checks=[]
def check(name,condition,detail=''):checks.append((name,bool(condition),str(detail)))
def norm(t):return re.sub(r'\s+','',t)
md=(BASE/'内容完整版.md').read_text(encoding='utf-8')
guide=(BASE/'阅读说明.md').read_text(encoding='utf-8')
worksheet=(BASE/'Tutorial1_答案参考.md').read_text(encoding='utf-8')
docx=BASE/'AI辅助肝癌候选分子发现内容手册.docx'
pdf=BASE/'AI辅助肝癌候选分子发现内容手册.pdf'
d=Document(docx);r=PdfReader(pdf)
dt=''.join(d.element.xpath('.//w:t/text()'))
pt='\n'.join(p.extract_text() or '' for p in r.pages)
check('Word重新打开',len(d.tables)==8 and len(d.inline_shapes)==4,f'tables={len(d.tables)} figures={len(d.inline_shapes)}')
check('PDF重新打开并确认页数',len(r.pages)==17,f'pages={len(r.pages)}')
check('首张解释表在第一页','白话解释表' in (r.pages[0].extract_text() or '') and '候选分子' in (r.pages[0].extract_text() or ''))
check('专业参数从附录开始','Kd' not in '\n'.join(p.extract_text() or '' for p in r.pages[:7]) and 'Kd' in (r.pages[7].extract_text() or ''))
check('Word标题样式及无装饰线',d.paragraphs[0].style.name=='Title' and not d.styles['Title']._element.xpath('.//w:pBdr'))
check('标题和标题层级均为黑色',all(d.styles[x].font.color.rgb and str(d.styles[x].font.color.rgb)=='000000' for x in ['Title','Subtitle','Heading 1','Heading 2']))
check('表格无固定行高及完整边框',all(not tr.xpath('./w:trPr/w:trHeight') for t in d.tables for tr in t._tbl.tr_lst) and all(c._tc.xpath('./w:tcPr/w:tcBorders') for t in d.tables for row in t.rows for c in row.cells))
check('表头重复标记',all(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for t in d.tables))
check('无空白PDF页',all(len(p.extract_text() or '')>100 for p in r.pages))
check('WordPDF内容主结论一致',all(norm(x) in norm(dt) and norm(x) in norm(pt) for x in ['提速幅度','未报告这些候选物的动物或患者试验','不是正常肝细胞','只是早期成果']))

md_claims=[]
for line in md.splitlines():
    t=line.strip()
    if not t or t.startswith('<!--') or t.startswith('```') or t.startswith('|---'):continue
    t=re.sub(r'^#{1,3} ','',t).replace('**','')
    t=re.sub(r'\[([^\]]+)\]\(https?://[^)]+\)',r'\1',t)
    if t.startswith('|'):md_claims.extend(x.strip() for x in t.strip('|').split('|'))
    else:md_claims.append(t)
missing_d=[x for x in md_claims if norm(x) not in norm(dt)]
missing_p=[x for x in md_claims if norm(x) not in norm(pt)]
check('源文档全部正文和表格写入Word',not missing_d,f'claims={len(md_claims)} missing={missing_d[:3]}')
check('源文档全部正文和表格写入PDF',not missing_p,f'claims={len(md_claims)} missing={missing_p[:3]}')

links=[]
for page in r.pages:
    for a in page.get('/Annots',[]):
        action=a.get_object().get('/A',{})
        if action.get('/URI'):links.append(str(action['/URI']))
expected_urls=set(re.findall(r'\]\((https?://[^)]+)\)',md))
check('全部文献链接可点击',expected_urls.issubset(set(links)),f'unique_expected={len(expected_urls)} pdf_links={len(links)}')

with (BASE/'证据清单.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
ids={row['证据编号'] for row in rows}
check('证据卡35条',len(rows)==35 and ids=={f'E{x:02d}' for x in range(1,36)})
refs=set(re.findall(r'\bE\d{2}\b',md))
for a,b in re.findall(r'E(\d{2})[–-]E(\d{2})',md):refs.update(f'E{x:02d}' for x in range(int(a),int(b)+1))
check('正文证据编号包括区间均有效',refs.issubset(ids),f'referenced={len(refs)}')
check('正文来源编号齐全',set(re.findall(r'\bS\d{2}\b',md)).issubset({f'S{x:02d}' for x in range(1,11)}))
for needle in ['9.2 ± 0.5 μM','566.7 ± 256.2 nM','360.0 ± 14.1 nM','33.4 ± 22.6 nM','208.7 ± 3.3 nM','1,706.7 ± 670.0 nM','IC50 > 6,000 nM']:
    check('WordPDF关键数值 '+needle,norm(needle) in norm(dt) and norm(needle) in norm(pt))
check('效率局限及两轮分界','没有同任务比较' in md and '第二轮结果与首轮30天终点要分开' in md and '没有把演练' in md)
check('旧靶点动物研究不移植为分子结果','不能改写成2023年分子的动物验证' in md)
check('日常生活联系','怎样连接日常生活' in md and '判断信息' in md and '治疗研究机会' in md)
check('rubric所有权重正确',all(x in md for x in ['Content 30%','Structure 20%','Literature 10%','Audience 10%','Delivery 10%','Visual Aids 10%','Critical Thinking 5%','Teamwork 5%']))
check('没有分配时间和成员',not re.search(r'(?m)^\s*\|?\s*(?:[01]?\d|20):[0-5]\d\s*[–-]',md) and not re.search(r'成员[一二三四1234]|每人\d+分钟',md))
check('未称真实受众已通过','尚未完成实际读者测试' in md)
check('Tutorial九题完整',all(re.search(r'(?m)^## '+str(i)+r' ',worksheet) for i in range(1,10)))
check('Tutorial书籍编辑正确','Vinícius Gonçalves Maltarollo' in worksheet and 'Victor G.' not in worksheet)
check('摘要和目录阅读层级透明','目前仅公开摘要和元数据' in md and '不计为章节已精读' in md)

book=PdfReader(BASE/'sources/Wiley_book_sample_3527312781.pdf')
btext=book.pages[6].extract_text() or ''
check('实际书章节重新打开',len(book.pages)==24 and 'Question 1' in btext and 'Question 2' in btext and 'selectivity' in btext)

asset_info=json.loads((BASE/'内容图示/图示说明.json').read_text(encoding='utf-8'))
for item in asset_info:
    png=BASE/'内容图示'/(item['name']+'.png');svg=BASE/'内容图示'/(item['name']+'.svg')
    with Image.open(png) as im:dimensions=im.size;im.verify()
    x=ET.parse(svg).getroot();texts=[e.text or '' for e in x.findall('{http://www.w3.org/2000/svg}text')]
    check('图示可重新打开 '+item['name'],dimensions==(1800,900) and item['labels']==texts)
check('四图都有说明且图四标明患者未报告',len(asset_info)==4 and '本案例未报告' in asset_info[3]['labels'])
preview=BASE/'_build/完整手册预览'
check('PDF每页渲染PNG齐全',all((preview/f'page-{i}.png').exists() for i in range(1,18)))
geometry=[]
with pdfplumber.open(pdf) as pp:
    for i,p in enumerate(pp.pages,1):
        bad=[c for c in p.chars if c['x0']<15 or c['x1']>p.width-15 or c['top']<12 or c['bottom']>p.height-12]
        if bad:geometry.append((i,len(bad)))
check('PDF文字无页面边缘溢出',not geometry,geometry)

visual=json.loads((BASE/'_build/complete_visual_qa.json').read_text(encoding='utf-8'))
check('最新WordPDF逐页人工视觉检查',visual.get('pages_reviewed')==list(range(1,18)) and visual.get('result')=='PASS' and visual.get('pdf_sha256')==hashlib.sha256(pdf.read_bytes()).hexdigest())
brief=json.loads((BASE/'_build/brief_checks.json').read_text(encoding='utf-8'))
check('原一页组员讨论稿保持不变',hashlib.sha256((BASE/'组员讨论用_研究问题微调.pdf').read_bytes()).hexdigest()==brief['sha256'])
check('已有资料全检查仍通过','FAIL=0' in (BASE/'验证记录.txt').read_text(encoding='utf-8'))

failed=[x for x in checks if not x[1]]
summary=f'COMPLETE_PACK_CHECKS={len(checks)} PASS={len(checks)-len(failed)} FAIL={len(failed)}'
lines=['完整内容包验证记录','日期 2026-10-03',summary,'Word与PDF均重新打开；Microsoft Word原生导出PDF，packaged render_docx.py用于栅格化。','']
lines += [f'{"PASS" if ok else "FAIL"} | {name} | {detail}' for name,ok,detail in checks]
lines += ['','执行命令',f'"C:\\Users\\woola\\.cache\\codex-runtimes\\codex-primary-runtime\\dependencies\\python\\python.exe" "{Path(__file__).resolve()}"','退出码 '+('1' if failed else '0'),'这些检查不等于同行评审、真实受众测试或评分保证。']
(BASE/'完整包验证记录.txt').write_text('\n'.join(lines),encoding='utf-8')
print(summary)
for name,ok,detail in failed:print('FAIL',name,detail)
print('REOPENED',docx);print('REOPENED',pdf)
sys.exit(1 if failed else 0)
