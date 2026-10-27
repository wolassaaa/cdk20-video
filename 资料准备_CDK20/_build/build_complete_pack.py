from pathlib import Path
from html import escape
import json, re, hashlib
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE=Path(__file__).resolve().parent.parent
ASSETS=BASE/'内容图示';ASSETS.mkdir(exist_ok=True)
OUT=BASE/'AI辅助肝癌候选分子发现内容手册.docx'
FONT='C:/Windows/Fonts/msyh.ttc'
BOLD='C:/Windows/Fonts/msyhbd.ttc'
BLUE='#254F6B';INK='#17242F';LIGHT='#F0F5F8';GRAY='#60707B'

class Diagram:
    def __init__(self,title):
        self.im=Image.new('RGB',(1800,900),'white');self.d=ImageDraw.Draw(self.im)
        self.svg=['<svg xmlns="http://www.w3.org/2000/svg" width="1800" height="900" viewBox="0 0 1800 900">', '<rect width="1800" height="900" fill="white"/>']
        self.labels=[];self.text(65,62,title,49,bold=True)
    def text(self,x,y,t,size=32,bold=False,color=INK):
        font=ImageFont.truetype(BOLD if bold else FONT,size)
        self.d.text((x,y),t,font=font,fill=color)
        self.svg.append(f'<text x="{x}" y="{y+size}" fill="{color}" font-family="Microsoft YaHei, sans-serif" font-size="{size}" font-weight="{700 if bold else 400}">{escape(t)}</text>')
        self.labels.append(t)
    def box(self,x,y,w,h,title,lines,status=True):
        fill=LIGHT if status else '#F5F5F5';border=BLUE if status else '#9AA2A8'
        self.d.rounded_rectangle((x,y,x+w,y+h),radius=14,fill=fill,outline=border,width=3)
        dash='' if status else ' stroke-dasharray="12 8"'
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="14" fill="{fill}" stroke="{border}" stroke-width="3"{dash}/>')
        self.text(x+24,y+25,title,35,True)
        for i,t in enumerate(lines):self.text(x+24,y+92+i*44,t,29)
    def line(self,points,arrow=True,color=BLUE):
        self.d.line(points,fill=color,width=5,joint='curve')
        self.svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in points)+f'" fill="none" stroke="{color}" stroke-width="5"/>')
        if arrow:
            x,y=points[-1];a,b=points[-2];dx=x-a;dy=y-b
            import math
            norm=math.hypot(dx,dy);ux,uy=dx/norm,dy/norm
            tri=[(x,y),(x-22*ux+11*uy,y-22*uy-11*ux),(x-22*ux-11*uy,y-22*uy+11*ux)]
            self.d.polygon(tri,fill=color)
            self.svg.append('<polygon points="'+' '.join(f'{a:.1f},{b:.1f}' for a,b in tri)+f'" fill="{color}"/>')
    def save(self,name,source):
        self.text(65,825,source,25,color=GRAY)
        self.svg.append('</svg>')
        (ASSETS/f'{name}.svg').write_text('\n'.join(self.svg),encoding='utf-8')
        self.im.save(ASSETS/f'{name}.png')
        return {'name':name,'labels':self.labels,'source':source,'dimensions':[1800,900]}

records=[]
g=Diagram('图一 我们到底在寻找什么')
g.box(70,245,480,250,'疾病研究需求',['肝细胞癌 HCC','我们想解决的研究问题'])
g.box(660,245,480,250,'药物作用对象',['CDK20 蛋白','潜在靶点 不是唯一病因'])
g.box(1250,245,480,250,'候选分子',['寻找值得测试的结构','早期研究对象 不是上市药'])
g.line([(550,370),(660,370)]);g.line([(1140,370),(1250,370)])
g.text(70,620,'主线：确定研究需求 → 选择作用对象 → 寻找并验证分子',38)
g.text(70,690,'聚焦一个靶点，不等于疾病只有一个原因。',33)
records.append(g.save('01_疾病靶点分子','依据 S01、S02、S10 | 教学概念图，不是实验图'))

g=Diagram('图二 AI 人和实验怎样合作')
xs=[70,680,1290]
top=[('1 疾病数据分析',['提出靶点优先级','研究者结合证据选择']),('2 结构预测',['AlphaFold 提供模型','研究者处理不确定区域']),('3 分子设计',['Chemistry42 提出结构','生成不等于合成'])]
bottom=[('4 计算评价',['对接等方法帮助排序','评分不等于真实效果']),('5 研究者选择',['考虑结构及合成等因素','确定优先实验对象']),('6 合成与实验',['真实制作和分层测试','实验结果用于继续设计'])]
for x,(title,lines) in zip(xs,top):g.box(x,190,440,200,title,lines)
for x,(title,lines) in zip(xs,bottom):g.box(x,505,440,200,title,lines)
g.line([(510,290),(680,290)]);g.line([(1120,290),(1290,290)])
g.line([(1510,390),(1510,445),(290,445),(290,505)])
g.line([(510,605),(680,605)]);g.line([(1120,605),(1290,605)])
g.text(70,751,'实验反馈 → 下一轮设计与测试   箭头表示步骤，不保证每一步成功。',32)
records.append(g.save('02_方法协作流程','依据 S02、S04、S05 | 原创流程解释，不是全自动系统'))

g=Diagram('图三 30天的终点在哪里')
g.box(70,190,440,190,'8,918 个',['生成结构','首轮计算探索'])
g.box(680,190,440,190,'54 个',['优先候选','不是实际合成数'])
g.box(1290,190,440,190,'7 个',['选择合成测试','不是7个全部成功'])
g.line([(510,285),(680,285)]);g.line([(1120,285),(1290,285)])
g.text(70,435,'30 天：靶点选择 → 首个实验确认的结合命中物',42,True)
g.text(70,500,'不是上市周期，也不是完整后续优化周期。',33)
g.box(70,595,770,145,'后续另一轮',['16 个新设计 → 6 个合成测试'])
g.box(960,595,770,145,'后续结果',['结合表现改善，并进一步做功能和细胞评价'])
g.line([(840,665),(960,665)])
records.append(g.save('03_首次命中与后续','依据 S02 E11、E13–E23 | 数量不是比例尺，不可据此计算提速倍数'))

g=Diagram('图四 这些证据回答不同问题')
labels=[('计算预测',['看起来有可能吗','研究提供预测'],True),('蛋白结合',['真的能结合吗','已有实验结果'],True),('蛋白功能',['能影响功能吗','已有实验结果'],True),('细胞实验',['细胞发生什么变化','已有模型结果'],True),('患者治疗',['患者是否获益','本案例未报告'],False)]
for i,(title,lines,ok) in enumerate(labels):g.box(65+i*350,250,290,260,title,lines,ok)
g.text(65,615,'结合、功能与细胞结果不能互相替代，更不能直接变成患者疗效。',36)
g.text(65,690,'2011年的靶点动物模型，不是2023年AI候选分子的动物验证。',31)
records.append(g.save('04_证据层级','依据 S02、S08 E14–E23、E33 | 四类报告结果，不表示疗效已经成立'))
(ASSETS/'图示说明.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')

def format_run(r,size=None,bold=None):
    r.font.name='Microsoft YaHei'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    r.font.color.rgb=RGBColor(0,0,0)
    if size:r.font.size=Pt(size)
    if bold is not None:r.bold=bold

def inline(p,text,size=None):
    pattern=r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))'
    for piece in re.split(pattern,text):
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',piece)
        if m:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
            fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Microsoft YaHei');fonts.set(qn('w:eastAsia'),'Microsoft YaHei');pr.append(fonts)
            color=OxmlElement('w:color');color.set(qn('w:val'),'254F6B');pr.append(color)
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int((size or 11)*2)));pr.append(sz)
            r.append(pr);t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
        else:
            bold=piece.startswith('**') and piece.endswith('**')
            format_run(p.add_run(piece[2:-2] if bold else piece),size,bold)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.7);sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Microsoft YaHei';s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    for e in list(s._element.findall('.//'+qn('w:pBdr'))):e.getparent().remove(e)
    fonts=s._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attr),None)
    s.paragraph_format.space_after=Pt(7)
    if name=='Normal':s.font.size=Pt(11);s.paragraph_format.line_spacing=1.2
    elif name=='Title':s.font.size=Pt(21);s.font.bold=True
    elif name=='Subtitle':s.font.size=Pt(9.5)
    else:s.font.size=Pt(15 if name=='Heading 1' else 12);s.font.bold=True;s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(10)
header=sec.header.paragraphs[0];header.text='SCIE1007  AI辅助肝癌候选分子发现';format_run(header.runs[0],8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
format_run(footer.add_run('内容准备版  |  '),8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='AI辅助肝癌候选分子发现内容手册'
doc.core_properties.subject='白话解释 完整内容 案例证据 来源评价 问答和rubric检查'
doc.core_properties.author='SCIE1007 项目内容准备'

def table(rows):
    n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=([1.3,2.9,2.9] if n==3 else [1.0,1.75,2.05,2.3])
    if '名称' in rows[0][0]:widths=[1.35,3.0,2.75]
    if '评分项' in rows[0][0]:widths=[1.15,3.3,2.65]
    if '资料类型' in rows[0][0]:widths=[1.5,1.9,3.7]
    for i,col in enumerate(t.columns):col.width=Inches(widths[i])
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for i,(c,txt) in enumerate(zip(cells,row)):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.1
            # Explicit wraps prevent Word's CJK compatibility settings from
            # letting punctuation-compressed runs touch adjacent cells.
            font=ImageFont.truetype(BOLD if ri==0 else FONT,102)
            capacity=(widths[i]*72-12)*10
            chunks=re.findall(r'[0-9]+(?:\.[0-9]+)?%|[A-Za-z0-9_]+|[^A-Za-z0-9_]',txt)
            wrapped=[];current=''
            for chunk in chunks:
                if current and font.getlength(current+chunk)>capacity:
                    wrapped.append(current.rstrip());current=chunk.lstrip()
                else:current+=chunk
            if current:wrapped.append(current.rstrip())
            inline(p,'\n'.join(wrapped),10.2)
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders)
            marg=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'60');el.set(qn('w:type'),'dxa');marg.append(el)
            pr.append(marg)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DBE7EE' if ri==0 else ('F7F9FA' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            if ri==0:
                for r in p.runs:r.bold=True
        trpr=t.rows[ri]._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def add_figures():
    for start in [0,2]:
        heading=doc.add_heading('解释图示 '+('一和二' if start==0 else '三和四'),level=1)
        heading.paragraph_format.page_break_before=True
        for item in records[start:start+2]:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
            p.add_run().add_picture(str(ASSETS/(item['name']+'.png')),width=Inches(7.05))
            cap=doc.add_paragraph();inline(cap,item['source'],9);cap.paragraph_format.space_after=Pt(12)

lines=(BASE/'内容完整版.md').read_text(encoding='utf-8').splitlines()
i=0;code=False;page_next=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- page -->':page_next=True;i+=1;continue
    if line.startswith('```'):code=not code;i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in row):rows.append(row)
            i+=1
        table(rows);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):
        if line[3:].startswith('附录六'):add_figures();page_next=True
        p=doc.add_heading(line[3:],level=1)
    elif line.startswith('### '):p=doc.add_heading(line[4:],level=2)
    else:
        p=doc.add_paragraph()
        if line.startswith('SCIE1007 '):p.style='Subtitle'
        inline(p,line)
        p.paragraph_format.widow_control=True
    if page_next:
        p.paragraph_format.page_break_before=True;page_next=False
    i+=1
doc.save(OUT)
reopened=Document(OUT)
print('DOCX_CREATED',OUT)
print('REOPENED paragraphs=',len(reopened.paragraphs),'tables=',len(reopened.tables),'diagrams=',len(reopened.inline_shapes))
print('DIAGRAMS_CREATED',len(records),'PNG and SVG pairs')
