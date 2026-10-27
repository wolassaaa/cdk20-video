from pathlib import Path
import csv,json,re,sys,hashlib,zipfile
from docx import Document
from pypdf import PdfReader
import pdfplumber
BASE=Path(__file__).resolve().parent.parent;OUT=BASE/'视频内容稿/教师要求补齐版';BACK=BASE/'_build/requirements_revision_20261007/before'
sys.path.insert(0,str(BASE/'_build/script_dependencies'))
from opencc import OpenCC
cc=OpenCC('s2t');tc=lambda s:cc.convert(s).replace('爲','為').replace('纔','才').replace('覈','核')
d=json.loads((OUT/'双语内容主稿.json').read_text(encoding='utf-8'));ps=[p for s in d['sections'] for p in s['paragraphs']];checks=[]
def check(n,v):checks.append((n,bool(v)))
norm=lambda s:re.sub(r'\s+','',s)
en=' '.join(p['en'] for p in ps);wc=len(re.findall(r"[A-Za-z0-9]+(?:[’'-][A-Za-z0-9]+)*",en))
check('8 chapters C01-C08; 38 continuous paragraph IDs',[s['id'] for s in d['sections']]==[f'C{i:02d}' for i in range(1,9)] and [p['id'] for p in ps]==[f'P{i:02d}' for i in range(1,39)])
check('2233 English words; estimate only',wc==2233==d['english_word_count'] and 15<wc/135<wc/120<20)
check('English-only voiceover matches canonical',(OUT/'英文口播纯文本.txt').read_text(encoding='utf-8')=='\n\n'.join(p['en'] for p in ps))
rows=list(csv.DictReader((OUT/'证据索引.csv').open(encoding='utf-8-sig')));oldrows=list(csv.DictReader((BACK/'证据索引.csv').open(encoding='utf-8-sig')))
check('Original 35 evidence rows unchanged; 2 explicitly conditional rows added',rows[:35]==oldrows and len(rows)==37 and [r['证据编号'] for r in rows[-2:]]==['E36','E37'])
ids={r['证据编号'] for r in rows}
check('All evidence IDs resolve',all(p['refs'] and set(p['refs'])<=ids for p in ps))
check('New benefits cite teacher guidance and conditional analysis',ps[32]['refs']==['E36'] and ps[33]['refs']==['E37'] and all(r['来源编号']=='T02;S02' and '不声称' in r['推论边界'] for r in rows[-2:]))
check('SC/TC numbers match',all(re.findall(r'\d+(?:,\d+)*',p['zh_sc'])==re.findall(r'\d+(?:,\d+)*',p['zh_tc']) for p in ps))
check('SC/TC reviewed conversion matches',all(tc(p['zh_sc'])==p['zh_tc'] for p in ps))
for name,phrase in [('opening','Can we predict whether a molecule is worth testing before we even make it?'),('counts','generated 8,918 structures'),('54 and 7','prioritised 54, then selected seven'),('30-day endpoint','30 days from target selection to the first experimentally confirmed binding hit'),('later16and6','proposed 16 structures and synthesised six'),('cell viability','reduced Huh7 liver cancer cell viability'),('no patient inference','a cell result still is not a patient result'),('case had no animal/patient trials','did not report animal or patient trials'),('fair traditional comparison','Traditional discovery is not blind guessing'),('fair future comparison','compare methods on the same target'),('daily life','judge the news'),('source evaluation','Scholarly books and reviews'),('health/cost/welfare potential','possibilities, not measured outcomes of this case'),('electricity trade-off','computers also consume electricity'),('resource comparison','both the calculations and the laboratory work'),('no net benefit asserted','this case does not establish that AI is cheaper'),('explicit closing answer','We can now answer our opening question')]:check(name,phrase in en)
check('No animation/shooting instructions in spoken content',not any(x in en.lower() for x in ['camera pans','cut to','shot list','animation direction']))
original=json.loads((BACK/'双语内容主稿.json').read_text(encoding='utf-8'));op=[p for s in original['sections'] for p in s['paragraphs']]
mapping=list(csv.DictReader((OUT/'修订段落对应表.csv').open(encoding='utf-8-sig')))
lookup={p['id']:p for p in ps};changed={'P02','P30','P36'}
check('Unedited scientific paragraphs preserved',all(lookup[r['新段号']]['en']==next(p['en'] for p in op if p['id']==r['旧段号']) for r in mapping if r['旧段号'] not in changed|{'new'}))
english_versions=[];page_counts={};hashes={}
for edition,field in [('简中','zh_sc'),('繁中','zh_tc')]:
 path=OUT/f'视频内容稿_英文加{edition}.docx';doc=Document(path);texts=[p.text for p in doc.paragraphs];alltext='\n'.join(texts)
 spoken=[t for t in texts if t in {p['en'] for p in ps}];english_versions.append(spoken)
 check(edition+' Word 38 English paragraphs in order',spoken==[p['en'] for p in ps])
 check(edition+' Word Chinese complete',all(p[field] in alltext for p in ps))
 check(edition+' English main title and objectives',texts[0]=='Can AI help design a molecule worth testing?' and all(o['en'] in alltext for o in d['objectives']))
 check(edition+' TOC field and cached internal hyperlinks',bool(doc.element.xpath('.//w:instrText[contains(text(),"TOC")]')) and len(doc.element.xpath('.//w:hyperlink[@w:anchor]'))>=10)
 check(edition+' Heading labels distinguish chapter and sources',len([p for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'C\d{2}',p.text)])==8)
 check(edition+' Visible benefits/future/conclusion subheadings',all(x in alltext for x in ['Benefits and trade-offs','Challenges and future directions','Conclusion and daily-life relevance']))
 check(edition+' Team on-screen and oral requirement explicit',('四位成員均須出鏡且口頭講解' if edition=='繁中' else '四位成员均须出镜且口头讲解') in alltext)
 pdf=path.with_suffix('.pdf');reader=PdfReader(pdf);page_counts[edition]=len(reader.pages);pt='\n'.join(p.extract_text() or '' for p in reader.pages)
 check(edition+' PDF all bilingual spoken content complete',all(norm(p['en']) in norm(pt) and norm(p[field]) in norm(pt) for p in ps))
 check(edition+' PDF all13pages readable',len(reader.pages)==13 and all(len(p.extract_text() or '')>150 for p in reader.pages))
 links=[a.get_object() for pg in reader.pages for a in pg.get('/Annots',[])]
 check(edition+' PDF has clickable navigation and 8 research URLs',len(links)>=18 and len({str(a.get('/A',{}).get('/URI')) for a in links if a.get('/A',{}).get('/URI')})==8)
 toc_text=reader.pages[0].extract_text()
 # Check cached contents page numbers against actual chapter heading locations, excluding TOC itself.
 good=True
 for section in d['sections']:
  title=section['id']+' '+section['title_en'];actual=next((n for n,p in enumerate(reader.pages,1) if n>1 and norm(title) in norm(p.extract_text() or '')),None)
  line=next((line for line in toc_text.splitlines() if line.strip().startswith(section['id'])), '')
  match=re.search(r'(\d+)\s*$',line);good=good and bool(match) and int(match.group(1))==actual
 check(edition+' All TOC page numbers correct',good)
 bad=[]
 with pdfplumber.open(pdf) as p:
  for n,pg in enumerate(p.pages,1):
   if any(c['x0']<12 or c['x1']>pg.width-12 or c['top']<10 or c['bottom']>pg.height-10 for c in pg.chars):bad.append(n)
 check(edition+' PDF text stays within page edges',not bad)
 for p in [path,pdf]:hashes[p.name]=hashlib.sha256(p.read_bytes()).hexdigest()
check('Both English versions identical',english_versions[0]==english_versions[1])
for name in ['视频内容稿_英文加简中.docx','视频内容稿_英文加繁中.docx','视频内容稿_英文加简中.pdf','视频内容稿_英文加繁中.pdf','双语内容主稿.json','证据索引.csv']:
 check('Original retained '+name,hashlib.sha256((BACK/name).read_bytes()).hexdigest()==hashlib.sha256((BASE/'视频内容稿'/name).read_bytes()).hexdigest())
result=f'CHECKS={len(checks)} PASS={sum(v for _,v in checks)} FAIL={sum(not v for _,v in checks)}'
record={'result':result,'english_words':wc,'estimated_minutes_at_120_135_wpm':[round(wc/135,1),round(wc/120,1)],'checks':[{'name':n,'pass':v} for n,v in checks],'pages':page_counts,'artifact_sha256':hashes,'actual_timed_reading':'PENDING','finished_video_and_four_presenters':'PENDING','new_teacher_announcements':'NOT_CHECKED'}
(BASE/'_build/requirements_revision_20261007/verification.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
print(result)
for n,v in checks:
 if not v:print('FAIL',n)
if any(not v for _,v in checks):sys.exit(1)
