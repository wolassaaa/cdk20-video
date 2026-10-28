"""Build the five Word files for the team from script v3.3: one master script and one per speaker (A-D),
every line in English with its Chinese translation, plus delivery (tone) and pace notes.
Source of truth: script_v3.3.md (English) and script_v3.3_zh.md (Chinese). Output: 分稿/*.docx
"""
import re, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
import sys; sys.path.insert(0, HERE)
from oncam_direction_v33 import DIRECTION, SHOOT, BROLL, RECORD, PROPS, TOGETHER
from oncam_graphics_v33 import GRAPHICS
OUT = os.path.join(HERE, "分稿")
os.makedirs(OUT, exist_ok=True)
en = open(os.path.join(HERE, "script_v3.3.md"), encoding="utf-8").read()
zh = open(os.path.join(HERE, "script_v3.3_zh.md"), encoding="utf-8").read()
en = en[en.index("## Part 3"):]
en = en[:en.index("## Part 4")] if "## Part 4" in en else en
HEAD = re.compile(r"^\*\*((?:\d-\d\d)|W\d) · ([\d:–]+) · (.+?)\*\*", re.M)


def parse(text, english):
    ms = list(HEAD.finditer(text))
    out = {}
    for k, m in enumerate(ms):
        body = text[m.end(): ms[k + 1].start() if k + 1 < len(ms) else len(text)]
        body = re.split(r"\n#{2,3} ", body)[0]
        if english:
            line = " ".join(l[2:].strip() for l in body.splitlines() if l.startswith("> "))
        else:
            paras = [p.strip() for p in body.strip().split("\n\n") if p.strip() and not p.strip().startswith("#")]
            line = paras[0] if paras else ""
        out[m.group(1)] = dict(id=m.group(1), time=m.group(2), head=m.group(3), text=line, pos=m.start())
    return out


EN, ZH = parse(en, True), parse(zh, False)
order = sorted(EN, key=lambda b: EN[b]["pos"])
secs = [(m.start(), m.group(1)) for m in re.finditer(r"^### (S\d · .+)$", en, re.M)]
section = lambda p: [n for s, n in secs if s <= p][-1]

SEC_ZH = {"S1": "开场", "S2": "研究问题", "S3": "第一个危机：没人见过这把锁", "S4": "第一个转折：AI 打开了灯",
          "S5": "第二个危机到第二个转折：钥匙转不动 → 钥匙转动了", "S6": "成绩单：以前 vs 现在", "S7": "下一步", "S8": "答案"}


def secs_to(t):
    m, s = t.split(":")
    return int(m) * 60 + int(s)


def spoken(b):
    return not b.startswith("W") and EN[b]["text"]


def who(b):
    m = re.match(r"([ABCD])\b", EN[b]["head"])
    return m.group(1) if m else None


def mode(b):
    return "出镜" if "ON-CAM" in EN[b]["head"] else "画外音"


# delivery notes for the lines that carry the story; everything else gets rule-based hints
KEY = {
    "The key went into the lock. But it barely turned.": "全片最低的一刻。放慢，声音往下沉，"
        "\"barely turned\" 前停半拍，说完不要急着收，让失望停在空气里。",
    "This time, the key turned.": "全片最高的一刻。说之前先停一拍（音乐会在这里空一拍），然后干脆、肯定地说出来，"
        "嘴角可以带一点笑。说完就是 8 秒纯音乐，不要接话。",
    "AI didn't replace the scientists. It ran their trial and error inside a computer, at light speed.":
        "全片的核心句。前半句平稳，\"at light speed\" 加重、稍快，像在揭晓答案。",
    "A docking score is a clue, not proof. Only the lab could tell.": "冷静、清醒，像给观众泼一点冷水，但不是泄气。",
}


def hints(b):
    t = EN[b]["text"]
    for k, v in KEY.items():
        if k in t:
            return v
    h = []
    if mode(b) == "出镜":
        h.append("看镜头，像在对一个朋友说话；比画外音再慢一点")
    if "?" in t:
        h.append("问句结尾上扬，问完停一拍")
    if re.search(r"\d", t):
        h.append("数字放慢、每个字咬清楚")
    if re.search(r"\b(But|And yet|However)\b", t):
        h.append("\"But\" 前面停一下，是转折")
    if "honest" in t.lower() or "not yet in people" in t or "Experiments still decide" in t:
        h.append("平静、真诚，不要像在道歉")
    return "；".join(h) or "自然叙述，句号处停顿"


def pace(b):
    """target speaking time at the film's base pace: 130 words a minute, on-camera lines a little slower (115)"""
    words = len(EN[b]["text"].split())
    wpm = 115 if mode(b) == "出镜" else 130
    return words, max(2, round(words / wpm * 60)), wpm


ROLE = {
    "A": dict(name="A · 开场人和向导",
              role="开场抛出问题、讲研究问题和计划、讲\"科学家的试错被搬进电脑\"、讲对接，最后讲下一步和全片最后一句。",
              tone="好奇、有带入感，像一个领着观众往前走的向导。开场要有钩子：先讲肝癌有多严重，再抛出\"如果电脑能走得更快呢？\"。"
                   "讲到 AI 那段要有一点兴奋，但不夸张。最后一句收得稳、干净。",
              speed="基准每分钟约 130 个英文单词（每秒 2 个多一点）。开场两句放慢到约 115 词/分钟，给观众时间进入状态。"),
    "B": dict(name="B · AI 和方法讲解人",
              role="讲 AlphaFold 怎么\"打开灯\"、尾巴为什么要剪掉、生成式设计的规则，以及成绩单里的部分对比。",
              tone="冷静、清楚、可信，像一个懂技术但讲人话的人。讲到 AlphaFold 预测出结构时，带一点惊叹；讲到\"科学家剪掉尾巴\"时，"
                   "语气里要有\"人还是要把关\"的分量。技术词（AlphaFold、Chemistry42、hinge）说得慢一点、清楚一点。",
              speed="基准约 130 词/分钟；技术名词和数字处放慢，整体比 A 稍稳。不要赶，宁可少说半拍。"),
    "C": dict(name="C · 结果主讲人（女声）",
              role="全片的情绪中心：讲 001 脱颖而出、\"钥匙几乎转不动\"、第二轮\"这一次钥匙转动了\"、细胞实验，以及最后回答研究问题。",
              tone="情绪起伏最大的一位。先是期待（七个里有一个脱颖而出），然后失落（几乎转不动），再是释放（转动了），"
                   "最后的答案要坚定、有底气。诚实那一句（在细胞里有效，还没在人身上）要平静真诚，不是道歉。",
              speed="基准约 130 词/分钟，但要跟着情绪变：失落那句放慢到约 100，转折那句干脆利落，结论部分中速、坚定。"),
    "D": dict(name="D · 概念和判断讲解人",
              role="讲锁和钥匙的比喻、为什么是 CDK20、\"在黑暗里怎么配钥匙\"，以及成绩单里\"以前 vs 现在\"的对比。",
              tone="沉稳、有画面感，像在讲一个故事。比喻部分要慢慢铺开，让观众\"看见\"那把锁；对比部分要有力度，"
                   "\"几百万个 → 13 个\"\"好几年 → 30 天\"这种数字对比，前后要停顿，让反差落地。",
              speed="基准约 130 词/分钟；比喻段落可以稍慢（约 120），对比的数字前后各停半拍。"),
}

GENERAL = [
    "语速基准：画外音每分钟约 130 个英文单词（每秒 2 个多一点），出镜句稍慢，约 115。表格里每句标了「约几秒」，录的时候差 10% 以内都没问题，最后以真实录音重新卡时间。",
    "出镜句：坐着，对着镜头，像在跟坐在对面的朋友聊天。每句拍 3 遍，选最自然的那遍。出镜句都在 19 个词以内，可以背下来。",
    "画外音：在安静的房间，用同一个麦克风录，嘴离麦克风约一拳远。",
    "停顿比速度更重要：句号停一下，问号后停一拍，\"But\" 前面停一下。",
    "情绪：这是一个有起伏的故事，不是读报告。前面有悬念，中间有失落和翻转，最后有答案。",
    "发音：CDK20 读 C-D-K-twenty · 001 读 zero-zero-one · 048 读 zero-four-eight · Kd 读 K-D · IC50 读 I-C-fifty · "
    "micromolar 读 MY-kro-MO-lar · nanomolar 读 NAN-o-MO-lar · Huh7 读 huh-seven · HEK293 读 heck two-nine-three · "
    "Nexavar 读 NEX-a-var · atezolizumab 读 a-te-zo-LIZ-yoo-mab · bevacizumab 读 bev-a-SIZ-yoo-mab · "
    "Chemistry42 读 chemistry forty-two · Insilico 读 in-SIL-i-co · lirafugratinib 读 leer-a-foo-GRA-ti-nib。",
]


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.6)
    return d


def heading(d, text, level):
    h = d.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.name = "Calibri"   # creates rFonts, so the East Asian face can be set
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return h


def para(d, text, bold=False, size=None, color=None, italic=False):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p


def beat_table(d, ids, member=None):
    cols = ["编号 · 时间", "谁 · 方式", "台词（英文）/ 中文意思", "语气 / 语速"]
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for b in ids:
        row = t.add_row().cells
        if b.startswith("W"):
            row[0].text = f"{b}\n{EN[b]['time']}"
            row[1].text = "纯音乐"
            row[2].text = f"（无人声）{ZH.get(b, {}).get('text', '')}"
            row[3].text = "不说话，让音乐和画面走"
            continue
        w, dur, wpm = pace(b)
        row[0].text = f"{b}\n{EN[b]['time']}"
        row[1].text = f"{who(b)} · {mode(b)}"
        cell = row[2]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(EN[b]["text"])
        r.bold = (member is not None and who(b) == member) or mode(b) == "出镜"
        r.font.size = Pt(11)
        p2 = cell.add_paragraph()
        r2 = p2.add_run(ZH.get(b, {}).get("text", ""))
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(90, 90, 90)
        row[3].text = f"{hints(b)}\n约 {dur} 秒（{w} 词，按每分钟 {wpm} 词）"
    widths = [Cm(2.2), Cm(2.0), Cm(9.6), Cm(4.6)]
    for r in t.rows:
        for i, wdt in enumerate(widths):
            r.cells[i].width = wdt


def guides(d, member=None):
    heading(d, "拍摄说明（先拍画面）", 1)
    para(d, "拍摄和录音分开做：这次只管\"演\"，声音之后单独录。", size=10, italic=True)
    for g in SHOOT:
        d.add_paragraph(g, style="List Bullet")
    heading(d, "道具（不用买）", 1)
    for g in PROPS:
        d.add_paragraph(g, style="List Bullet")
    heading(d, "对话段：尽量坐在一起拍", 1)
    for g in TOGETHER:
        d.add_paragraph(g)
    heading(d, "空镜（每人拍一点，剪辑时垫用）", 1)
    for g in BROLL:
        d.add_paragraph(g, style="List Bullet")
    heading(d, "录音说明（之后单独录声音）", 1)
    for g in RECORD:
        d.add_paragraph(g, style="List Bullet")
    ids = [b for b in order if spoken(b) and mode(b) == "出镜" and (member is None or who(b) == member)]
    heading(d, "出镜句的动作指导", 1)
    para(d, "出镜时你的画面是一张暖色的“照片卡”，占半屏，另一半是黑底墨线图示（后期做，你不用管）。标【道具】的用手边的东西，标【后期】的是我画上去的。知道旁边会出现什么，可以帮你掌握说话的节奏和手势方向。",
         size=10, italic=True)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, c in enumerate(["编号 · 谁", "台词 / 中文意思", "动作 · 表情和眼神", "参考", "旁边的图示（后期）"]):
        t.rows[0].cells[i].text = c
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for b in ids:
        act, face, ref = DIRECTION.get(b, ("站定，直视镜头，自然说出来。", "自然。", "—"))
        gfx, when = GRAPHICS.get(b, ("—", ""))
        row = t.add_row().cells
        row[0].text = f"{b}\n{who(b)}"
        row[1].text = ""
        rr = row[1].paragraphs[0].add_run(EN[b]["text"]); rr.bold = True
        p2 = row[1].add_paragraph(); r2 = p2.add_run(ZH.get(b, {}).get("text", "")); r2.font.size = Pt(9.5); r2.font.color.rgb = RGBColor(90, 90, 90)
        row[2].text = f"动作：{act}\n表情和眼神：{face}"
        row[3].text = ref
        row[4].text = f"{gfx}\n出现时机：{when}"
    for rw in t.rows:
        for i, wdt in enumerate([Cm(1.5), Cm(4.6), Cm(5.0), Cm(2.8), Cm(4.8)]):
            rw.cells[i].width = wdt


# ---------- master ----------
d = new_doc()
heading(d, "WORTH TESTING? · 总稿 v3.3（中英对照）", 0)
para(d, "SCIE1007 Topic 10 · 研究问题：How can computer-aided methods help scientists find and assess potential cancer drug candidates? "
        "A case study of CDK20 in liver cancer.", italic=True)
para(d, "全片约 17 分 43 秒（含纯音乐段）· A 约 3:47 · B 约 3:49 · C 约 4:17 · D 约 3:59。粗体 = 出镜句。", size=10)
heading(d, "所有人都要读的录音要求", 1)
for g in GENERAL:
    d.add_paragraph(g, style="List Bullet")
guides(d)
heading(d, "四个人的分工", 1)
for k, r in ROLE.items():
    para(d, r["name"], bold=True)
    d.add_paragraph(r["role"])
cur = None
group = []
for b in order:
    s = section(EN[b]["pos"])
    if s != cur:
        if group:
            beat_table(d, group)
            group = []
        cur = s
        heading(d, f"{s}　｜　{SEC_ZH.get(s.split(' ')[0], '')}", 1)
    group.append(b)
if group:
    beat_table(d, group)
d.save(os.path.join(OUT, "00_总稿_v3.3_中英对照.docx"))

# ---------- one per member ----------
spoken_order = [b for b in order if spoken(b)]
for m, r in ROLE.items():
    mine = [b for b in spoken_order if who(b) == m]
    words = sum(len(EN[b]["text"].split()) for b in mine)
    oncam = [b for b in mine if mode(b) == "出镜"]
    d = new_doc()
    heading(d, f"WORTH TESTING? · {r['name']} · 分稿 v3.3", 0)
    para(d, f"你一共 {len(mine)} 句，{words} 个英文单词，约 {words / 130:.1f} 分钟；其中出镜 {len(oncam)} 句（要背、要对镜头），"
            f"其余是画外音（录音即可）。", size=11)
    heading(d, "你的角色", 1)
    d.add_paragraph(r["role"])
    heading(d, "语气要求", 1)
    d.add_paragraph(r["tone"])
    heading(d, "语速要求", 1)
    d.add_paragraph(r["speed"])
    heading(d, "通用录音要求", 1)
    for g in GENERAL:
        d.add_paragraph(g, style="List Bullet")
    guides(d, m)
    heading(d, "要背下来的出镜句", 1)
    for b in oncam:
        p = para(d, f"{b}　{EN[b]['text']}", bold=True, size=11.5)
        para(d, f"　　{ZH.get(b, {}).get('text', '')}", size=10, color=(90, 90, 90))
    heading(d, "你的全部台词（按出场顺序）", 1)
    para(d, "「接在谁后面」是你开口前一句话的结尾，听到它就该你了。", size=10, italic=True)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, c in enumerate(["编号 · 时间 · 方式", "接在谁后面", "台词（英文）/ 中文意思", "语气 / 语速"]):
        t.rows[0].cells[i].text = c
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for b in mine:
        i = order.index(b)
        prev = next((x for x in reversed(order[:i]) if spoken(x) or x.startswith("W")), None)
        if prev is None:
            cue = "（全片开头）"
        elif prev.startswith("W"):
            cue = f"（纯音乐 {prev} 结束后）"
        else:
            cue = f"{who(prev)}：…" + " ".join(EN[prev]["text"].split()[-6:])
        w, dur, wpm = pace(b)
        row = t.add_row().cells
        row[0].text = f"{b}\n{EN[b]['time']}\n{mode(b)}"
        row[1].text = cue
        cell = row[2]
        cell.text = ""
        rr = cell.paragraphs[0].add_run(EN[b]["text"])
        rr.bold = True
        rr.font.size = Pt(11.5)
        p2 = cell.add_paragraph()
        r2 = p2.add_run(ZH.get(b, {}).get("text", ""))
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(90, 90, 90)
        row[3].text = f"{hints(b)}\n约 {dur} 秒（{w} 词，按每分钟 {wpm} 词）"
    for rw in t.rows:
        for i, wdt in enumerate([Cm(2.4), Cm(3.4), Cm(8.4), Cm(4.2)]):
            rw.cells[i].width = wdt
    d.save(os.path.join(OUT, f"0{'ABCD'.index(m) + 1}_{m}_分稿_v3.3_中英对照.docx"))
    print(m, len(mine), "lines", words, "words", len(oncam), "on-cam")
print("files:", sorted(os.listdir(OUT)))
